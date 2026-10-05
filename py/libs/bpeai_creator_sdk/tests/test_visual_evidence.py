from __future__ import annotations

import base64
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from PIL import Image
from pptx import Presentation

from bpeai_creator_sdk.artifacts.hero_image import (
    attach_title_hero_image, engineering_basis_sha256, engineering_summary,
    load_visual_catalog, select_verified_visual,
)
from bpeai_creator_sdk.artifacts.pptx_eval import build_evaluation_pptx, build_slide_pack_from_evaluation
from bpeai_creator_sdk.sme.pack_loader import knowledge_pack_from_dict, load_knowledge_pack


@pytest.fixture
def visual(tmp_path):
    path = tmp_path / "references/visuals/pump.png"
    path.parent.mkdir(parents=True)
    # A synthetic fixture only; never shipped as vendor evidence.
    Image.new("RGB", (800, 400), "#17324D").save(path)
    data = path.read_bytes()
    result = {"schema_version": "equipment_selector_v1", "system_name": "Transfer pump",
              "equipment_name": "Transfer pump", "selected_model": "Centrifugal pump",
              "rationale": "Suitable for the supplied transfer duty", "dir_code": "1-2-3"}
    asset = {"id": "pump", "path": "references/visuals/pump.png", "kind": "vendor_photo",
             "scope": "product_family", "manufacturer": "Test vendor", "model": "Test family",
             "match": {"equipment_name": "Transfer pump", "selected_model": "Centrifugal pump"},
             "source_url": "https://example.org/catalog", "attribution": "Test vendor",
             "review_status": "approved", "reviewed_by": "Fixture reviewer", "reviewed_at": "2026-10-05",
             "applicability": "Example of the recommended pump technology",
             "reuse_basis": "Synthetic test fixture", "sha256": hashlib.sha256(data).hexdigest()}
    catalog = {"schema_version": "ei_visual_assets_v1", "assets": [asset]}
    (tmp_path / "references/visuals/catalog.json").write_text(json.dumps(catalog), encoding="utf-8")
    return result, catalog, tmp_path, data


def test_local_and_database_payload_produce_same_visual(visual):
    result, catalog, root, data = visual
    for name in ("pack", "dir_requirements", "equipment_options", "validation_rules"):
        (root / f"{name}.yaml").write_text("pack_id: fixture\n" if name == "pack" else "{}\n")
    # Explicit pack_root is the parent; use the actual folder name for this fixture.
    local = load_knowledge_pack(root.name, pack_root=root.parent)
    payload = copy.deepcopy(catalog)
    payload["assets"][0]["data_base64"] = base64.b64encode(data).decode()
    remote = knowledge_pack_from_dict("fixture", {"content": {"visual_assets": payload}})
    assert select_verified_visual(result, local.visual_assets, local.path) == select_verified_visual(result, remote.visual_assets)
    local_slides, remote_slides = {}, {}
    a = attach_title_hero_image(result, local_slides, output_path=root / "local.png", knowledge_pack=local)
    b = attach_title_hero_image(result, remote_slides, output_path=root / "remote.png", knowledge_pack=remote)
    assert a.read_bytes() == b.read_bytes() == data
    assert local_slides["visual_evidence"] == remote_slides["visual_evidence"]
    assert "data_base64" not in remote_slides["visual_evidence"]


@pytest.mark.parametrize("changes", [
    {"review_status": "draft"}, {"reviewed_by": ""}, {"reviewed_at": "invalid"},
    {"source_url": ""}, {"reuse_basis": ""}, {"sha256": "0" * 64},
    {"kind": "ai_generated"}, {"scope": "exact_model"},
    {"kind": []}, {"scope": {}}, {"reviewed_by": {"name": "invalid"}},
    {"match": {"equipment_name": "Different pump", "selected_model": "Centrifugal pump"}},
    {"match": {"equipment_type": "pump"}}, {"path": "../pump.png"},
    {"path": "C:/pump.png"}, {"path": "references/visuals/../../pump.png"},
])
def test_unverified_or_inapplicable_assets_fall_back(visual, changes):
    result, catalog, root, _ = visual
    catalog["assets"][0].update(changes)
    assert select_verified_visual(result, catalog, root) is None


def test_missing_or_corrupt_asset_falls_back(visual):
    result, catalog, root, _ = visual
    (root / catalog["assets"][0]["path"]).unlink()
    assert select_verified_visual(result, catalog, root) is None
    catalog["assets"][0]["data_base64"] = "not base64!"
    assert select_verified_visual(result, catalog) is None


def test_project_drawing_invalidated_when_engineering_basis_changes(visual):
    result, catalog, root, _ = visual
    asset = catalog["assets"][0]
    asset.update(kind="project_drawing", scope="project_configuration", result_sha256=engineering_basis_sha256(result))
    asset["match"]["dir_code"] = result["dir_code"]
    assert select_verified_visual(result, catalog, root)
    result["dimensions"] = {"value": "2", "unit": "m"}
    assert select_verified_visual(result, catalog, root) is None


def test_summary_rejects_legacy_and_llm_images(visual):
    result, catalog, root, _ = visual
    legacy = str(root / catalog["assets"][0]["path"])
    result["hero_image_path"] = legacy
    result["artifacts"] = {"hero_image_path": legacy}
    slides = build_slide_pack_from_evaluation(result)
    slides.update(hero_image_path=legacy, visual_evidence={"mode": "asset", **catalog["assets"][0]})
    slides["slides"][0].update(hero_image_path=legacy, hero_tags=["GMP ready"])
    path = build_evaluation_pptx(result, slide_pack=slides, output_path=root / "summary.pptx")
    title = Presentation(path).slides[0]
    assert not any(shape.shape_type == 13 for shape in title.shapes)
    text = " ".join(shape.text for shape in title.shapes if shape.has_text_frame)
    assert "Engineering summary" in text and "Centrifugal pump" in text
    assert "GMP ready" not in text and "vendor available" not in text
    assert "hero_image_path" not in slides
    assert slides["visual_evidence"]["mode"] == "summary"


def test_verified_picture_keeps_proportions_caption_and_notes(visual):
    result, catalog, root, _ = visual
    pack = SimpleNamespace(path=root, visual_assets=catalog)
    slides = build_slide_pack_from_evaluation(result)
    path = build_evaluation_pptx(result, output_path=root / "asset.pptx", slide_pack=slides, knowledge_pack=pack)
    title = Presentation(path).slides[0]
    picture = next(shape for shape in title.shapes if shape.shape_type == 13)
    assert picture.width / picture.height == pytest.approx(2)
    text = " ".join(shape.text for shape in title.shapes if shape.has_text_frame)
    assert "Product-family example" in text and "Configuration is not confirmed" in text
    assert "Test vendor" in text
    assert catalog["assets"][0]["sha256"] in title.notes_slide.notes_text_frame.text
    assert slides["visual_evidence"]["mode"] == "asset"


def test_sizing_summary_preserves_basis_and_omits_unknowns():
    result = {"schema_version": "equipment_sizing_v1", "capacity": {"value": "10", "unit": "m3/h", "basis": "DIR"},
              "dimensions": {"value": "unknown", "unit": "m"},
              "connections": [{"name": "Inlet", "size": "TBD", "unit": "in"}],
              "assumptions": ["Vendor to confirm seal configuration"]}
    rows = dict(engineering_summary(result))
    assert rows["Capacity / duty"] == "10 m3/h" and rows["Sizing basis"] == "DIR"
    assert "Envelope" not in rows and "Connections" not in rows
    assert "Vendor to confirm" in rows["Open items / assumptions"]


def test_malformed_catalog_is_optional(tmp_path):
    path = tmp_path / "references/visuals/catalog.json"
    path.parent.mkdir(parents=True)
    path.write_text("broken")
    assert load_visual_catalog(tmp_path) == {}


def test_cli_bundle_contains_catalog_and_binary(visual):
    _, _, root, data = visual
    py_root = Path(__file__).resolve().parents[3]
    spec = importlib.util.spec_from_file_location("upload_visual_fixture", py_root / "tools/upload_creator_bundle.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    import zipfile
    with zipfile.ZipFile(root / "bundle.zip", "w") as archive:
        module.add_tree(archive, root / "references", "py/knowledge/fixture/references")
    with zipfile.ZipFile(root / "bundle.zip") as archive:
        assert archive.read("py/knowledge/fixture/references/visuals/pump.png") == data
        assert "py/knowledge/fixture/references/visuals/catalog.json" in archive.namelist()


@pytest.mark.parametrize("zip_download", [False, True])
def test_download_cli_restores_portable_visuals(visual, monkeypatch, zip_download):
    import io
    import sys
    import zipfile

    result, catalog, root, data = visual
    py_root = Path(__file__).resolve().parents[3]
    spec = importlib.util.spec_from_file_location("download_visual_fixture", py_root / "tools/download_knowledge_pack.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    relative = catalog["assets"][0]["path"]
    if zip_download:
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w") as archive:
            archive.writestr("fixture/references/visuals/catalog.json", json.dumps(catalog))
            archive.writestr("fixture/" + relative, data)
        response = io.BytesIO(stream.getvalue())
        response.headers = {"X-BPEAI-Content-Version": "1"}
    else:
        response = io.BytesIO(json.dumps({"pack_id": "fixture",
            "files": {"references/visuals/catalog.json": json.dumps(catalog)},
            "binary_files": {relative: base64.b64encode(data).decode()}}).encode())
    monkeypatch.setattr(module.urllib.request, "urlopen", lambda *args, **kwargs: response)
    monkeypatch.setenv("BPEAI_SESSION_COOKIE", "fixture-only")
    output = root / "download"
    monkeypatch.setattr(sys, "argv", ["download", "--pack", "fixture", "--out", str(output)] + (["--zip"] if zip_download else []))
    assert module.main() == 0
    restored = output / "fixture"
    assert (restored / relative).read_bytes() == data
    assert select_verified_visual(result, load_visual_catalog(restored), restored)[1] == data


@pytest.mark.parametrize("family", ["equipment_evaluator", "equipment_sizing"])
@pytest.mark.parametrize("runtime", ["local", "website"])
def test_template_pptx_phase_uses_loaded_pack_evidence(visual, monkeypatch, family, runtime):
    """Exercise the same run() handshake used by local_chat and the website UI."""
    import sys
    from bpeai_creator_sdk.local_run import load_agent_class

    result, catalog, root, data = visual
    py_root = Path(__file__).resolve().parents[3]
    # Another SDK import test may have cached a temporary 'apps' namespace.
    old_modules = {key: value for key, value in sys.modules.items() if key == "apps" or key.startswith("apps.")}
    for key in old_modules:
        monkeypatch.delitem(sys.modules, key)
    monkeypatch.syspath_prepend(str(py_root))
    cls = load_agent_class(family, py_root=py_root)
    monkeypatch.setattr(sys.modules[cls.__module__], "repo_py_root", lambda: py_root)
    agent = cls()
    if family == "equipment_sizing":
        result["schema_version"] = "equipment_sizing_v1"
        result["capacity"] = {"value": "12", "unit": "m3/h", "basis": "DIR"}
    payload = {"meta": {"pack_id": "fixture"}, "content": {"visual_assets": copy.deepcopy(catalog)}}
    payload["content"]["visual_assets"]["assets"][0]["data_base64"] = base64.b64encode(data).decode()
    local_pack = knowledge_pack_from_dict("fixture", {"content": {"visual_assets": catalog}}, path=root)
    monkeypatch.setattr(agent, "_ensure_knowledge_pack", lambda *a, **kw: (local_pack, []))
    monkeypatch.setattr(agent, "_resolve_or_generate_dir_menu", lambda *a, **kw: (SimpleNamespace(), []))
    monkeypatch.setattr(agent, "_build_pptx_slide_pack", lambda pack, value: build_slide_pack_from_evaluation(value))
    monkeypatch.chdir(root)
    inputs = {"phase": "pptx", "system_name": result["system_name"], "evaluation_result": result}
    if runtime == "website":
        inputs["knowledge_pack_payload"] = payload
    output = agent.run(inputs)
    assert output["artifacts"]["visual_evidence"]["mode"] == "asset"
    assert Path(output["artifacts"]["hero_image_path"]).read_bytes() == data
    assert output["deliverable"] == "pptx"
    json.dumps(output)  # No in-memory image objects / bytes leak into the wire result.
    for key in list(sys.modules):
        if key == "apps" or key.startswith("apps."):
            del sys.modules[key]
    sys.modules.update(old_modules)
