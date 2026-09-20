"""Sizing report DOCX — datasheet markdown, not option-evaluation layout."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping

from .docx_eval import (
    CALLOUT_BG,
    MUTED,
    NAVY,
    _add_callout,
    _add_paragraph,
    _add_table,
    _append_unstructured_markdown,
    _set_cell_text,
    _set_header_footer,
    _set_run,
    _shade,
)
from .report_content import _item_noun


def _sizing_item_noun(result: Mapping[str, Any]) -> str:
    item = str(result.get("sized_item") or "").strip()
    if item:
        return item
    return _item_noun(result)


def _display_title(result: Mapping[str, Any], title: str | None) -> str:
    if title and str(title).strip():
        return str(title).strip()
    try:
        from .names import sizing_display_title

        return sizing_display_title(result)
    except ImportError:
        return str(result.get("system_name") or result.get("equipment_name") or "Sizing")


def _key_spec_rows(result: Mapping[str, Any]) -> list[list[str]]:
    rows: list[list[str]] = []
    for spec in result.get("key_specs") or []:
        if not isinstance(spec, Mapping):
            continue
        key = str(spec.get("key") or "").strip()
        if not key:
            continue
        value = str(spec.get("value") or "").strip()
        unit = str(spec.get("unit") or "").strip()
        rows.append([key, value, unit])
    return rows


def write_sizing_report_docx(
    result: Mapping[str, Any],
    output_path: Path | str,
    *,
    title: str | None = None,
) -> Path:
    """Render equipment_sizing_v1 as an editable Word report."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    out = Path(output_path)
    if out.suffix.lower() != ".docx":
        out = out.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)

    markdown = (result.get("datasheet_markdown") or "").strip()
    if not markdown:
        cap = result.get("capacity") if isinstance(result.get("capacity"), Mapping) else {}
        markdown = (
            f"# {result.get('system_name') or 'Equipment'} sizing\n\n"
            f"## Capacity\n\n{cap.get('value', '')} {cap.get('unit', '')}\n"
        )
    display_title = _display_title(result, title)
    dir_code = str(result.get("dir_code") or "").strip()
    application = str(result.get("application") or "").strip()
    item_noun = _sizing_item_noun(result)
    subtitle = f"Preliminary {item_noun} sizing"
    header_title = display_title
    if dir_code:
        header_title += f" • Basis: user DIR code {dir_code}"

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.7)
    _set_header_footer(doc, header_title)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_p.add_run(display_title)
    _set_run(title_run, size_pt=20, bold=True, color=NAVY)
    title_p.paragraph_format.space_after = Pt(2)

    _add_paragraph(doc, subtitle, size_pt=10, color=MUTED, space_after=8)

    meta = [
        c
        for c in (
            f"Validated DIR: {dir_code}" if dir_code else "",
            f"Application: {application}" if application else "",
            f"Sized item: {item_noun}" if item_noun else "",
        )
        if c
    ]
    if meta:
        meta_table = doc.add_table(rows=1, cols=len(meta))
        meta_table.style = "Table Grid"
        for i, cell_text in enumerate(meta):
            _shade(meta_table.rows[0].cells[i], CALLOUT_BG)
            _set_cell_text(meta_table.rows[0].cells[i], cell_text)
            if meta_table.rows[0].cells[i].paragraphs[0].runs:
                _set_run(
                    meta_table.rows[0].cells[i].paragraphs[0].runs[0],
                    size_pt=9,
                    bold=True,
                    color=NAVY,
                )
        doc.add_paragraph("")

    selected = str(result.get("selected_model") or "").strip()
    if selected:
        _add_callout(doc, "Sized-item concept", selected)

    spec_rows = _key_spec_rows(result)
    if spec_rows:
        _add_paragraph(doc, "Key specifications", size_pt=13, bold=True, color=NAVY, space_before=8, space_after=4)
        _add_table(doc, ("Item", "Value", "Unit"), spec_rows)

    _append_unstructured_markdown(doc, markdown, str(result.get("system_name") or ""), display_title)

    try:
        doc.save(str(out))
        return out
    except PermissionError:
        from datetime import datetime

        stamped = out.with_name(f"{out.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{out.suffix}")
        doc.save(str(stamped))
        return stamped


def build_sizing_docx(
    result: Mapping[str, Any],
    *,
    output_path: Path | str,
    title: str | None = None,
) -> Path:
    """Public SDK entry for equipment_sizing Word reports."""
    return write_sizing_report_docx(result, output_path, title=title)


def sizing_docx_text(path: Path | str) -> str:
    """Extract paragraph and table text for tests."""
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)
