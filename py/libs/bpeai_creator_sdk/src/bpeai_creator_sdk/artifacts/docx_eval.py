"""Evaluation report DOCX matching the former URS-style PDF layout.

Content stays structured evaluate JSON + ``datasheet_markdown``. Layout is
Python via python-docx so users can edit, then Save as PDF.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Mapping, Sequence

from .report_content import (
    _as_str_list,
    _decoded_rows,
    _eval_options,
    _exclusion_rows,
    _extract_md_table,
    _extract_recommendation,
    _fit_position,
    _has_structured_eval,
    _heading_text,
    _is_table_divider,
    _item_noun,
    _key_spec_pairs,
    _labeled_bullet,
    _markdown_section,
    _objective_rows,
    _rating_dots,
    _section_intro,
    _split_labeled,
    _split_table_row,
    _synthesize_markdown,
)

NAVY = "1F4E79"
WHITE = "FFFFFF"
MUTED = "5B6770"
BODY = "222222"
HEADER_BG = "1F4E79"
ROW_ALT = "F5F8FC"
CALLOUT_BG = "E8EFF7"
GRID = "C5CDD4"

_HTML_TAG_RE = re.compile(r"<[^>]+>")
_MD_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _display_title(result: Mapping[str, Any], title: str | None) -> str:
    if title and str(title).strip():
        return str(title).strip()
    try:
        from .names import evaluation_display_title

        return evaluation_display_title(result)
    except ImportError:
        return str(result.get("system_name") or result.get("equipment_name") or "Evaluation")


def _plain(text: Any) -> str:
    s = str(text or "")
    s = s.replace("<br/>", "\n").replace("<br />", "\n").replace("<br>", "\n")
    s = _HTML_TAG_RE.sub("", s)
    s = _MD_BOLD_RE.sub(r"\1", s)
    s = s.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
    return s.strip()


def _rgb(hex_color: str):
    from docx.shared import RGBColor

    h = hex_color.strip().lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade(cell, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_run(run, *, size_pt: float, bold: bool = False, color: str = BODY) -> None:
    from docx.shared import Pt

    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = _rgb(color)
    run.font.name = "Calibri"


def _add_paragraph(
    container,
    text: str,
    *,
    size_pt: float = 10,
    bold: bool = False,
    color: str = BODY,
    space_before: float = 0,
    space_after: float = 6,
    label: str = "",
) -> None:
    from docx.shared import Pt

    para = container.add_paragraph()
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after = Pt(space_after)
    if label:
        lead = para.add_run(label)
        _set_run(lead, size_pt=size_pt, bold=True, color=NAVY)
        rest = para.add_run(_plain(text))
        _set_run(rest, size_pt=size_pt, bold=False, color=color)
        return
    run = para.add_run(_plain(text))
    _set_run(run, size_pt=size_pt, bold=bold, color=color)


def _set_cell_text(
    cell,
    text: str,
    *,
    header: bool = False,
    label_col: bool = False,
) -> None:
    from docx.shared import Pt

    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(2)
    para.paragraph_format.space_after = Pt(2)
    raw = _plain(text)
    if header:
        _shade(cell, HEADER_BG)
        run = para.add_run(raw)
        _set_run(run, size_pt=9, bold=True, color=WHITE)
        return
    if label_col:
        _shade(cell, HEADER_BG)
        run = para.add_run(raw)
        _set_run(run, size_pt=9, bold=True, color=WHITE)
        return
    for i, line in enumerate(raw.split("\n")):
        if i:
            para.add_run("\n")
        run = para.add_run(line)
        _set_run(run, size_pt=9, bold=False, color=BODY)


def _add_table(doc, headers: Sequence[str], rows: Sequence[Sequence[str]]) -> None:
    if not headers or not rows:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    for i, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], str(header), header=True)
    for r, row in enumerate(rows, start=1):
        padded = list(row) + [""] * len(headers)
        for c in range(len(headers)):
            _set_cell_text(table.rows[r].cells[c], str(padded[c]))
            if r % 2 == 0:
                _shade(table.rows[r].cells[c], ROW_ALT)
    doc.add_paragraph("")


def _add_label_table(doc, rows: Sequence[Sequence[str]]) -> None:
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        _set_cell_text(table.rows[r].cells[0], str(row[0] if row else ""), label_col=True)
        _set_cell_text(table.rows[r].cells[1], str(row[1] if len(row) > 1 else ""))
    doc.add_paragraph("")


def _add_callout(doc, title: str, body: str) -> None:
    table = doc.add_table(rows=2 if title else 1, cols=1)
    table.style = "Table Grid"
    if title:
        _shade(table.rows[0].cells[0], CALLOUT_BG)
        _set_cell_text(table.rows[0].cells[0], title)
        _shade(table.rows[0].cells[0], CALLOUT_BG)
        run = table.rows[0].cells[0].paragraphs[0].runs[0]
        _set_run(run, size_pt=9, bold=True, color=NAVY)
        _shade(table.rows[1].cells[0], CALLOUT_BG)
        _set_cell_text(table.rows[1].cells[0], body)
        _shade(table.rows[1].cells[0], CALLOUT_BG)
    else:
        _shade(table.rows[0].cells[0], CALLOUT_BG)
        _set_cell_text(table.rows[0].cells[0], body)
        _shade(table.rows[0].cells[0], CALLOUT_BG)
    doc.add_paragraph("")


def _add_md_table_or(
    doc,
    section: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
) -> None:
    parsed = _extract_md_table(section)
    if parsed:
        _add_table(doc, parsed[0], parsed[1])
        return
    if headers and rows:
        _add_table(doc, headers, rows)


def _set_header_footer(doc, header_title: str) -> None:
    section = doc.sections[0]
    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.text = ""
    run = hp.add_run(header_title[:160])
    _set_run(run, size_pt=8, color=MUTED)
    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.text = ""
    run = fp.add_run("Edit in Word, then Save as PDF if a locked copy is needed.")
    _set_run(run, size_pt=8, color=MUTED)


def _append_unstructured_markdown(doc, markdown: str, system: str, display_title: str) -> None:
    skip_titles = {system.lower(), display_title.lower(), "evaluation"}
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or (line.startswith("---") and not line.strip().startswith("|")):
            i += 1
            continue
        heading = _heading_text(line)
        if heading:
            lowered = heading.lower()
            if lowered in skip_titles or lowered.startswith("validated dir"):
                i += 1
                continue
            _add_paragraph(doc, heading, size_pt=12, bold=True, color=NAVY, space_before=10, space_after=4)
            i += 1
            continue
        if line.strip().startswith("|") and i + 1 < len(lines) and _is_table_divider(lines[i + 1]):
            headers = _split_table_row(line)
            rows: list[list[str]] = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _is_table_divider(lines[i]):
                    rows.append(_split_table_row(lines[i]))
                i += 1
            _add_table(doc, headers, rows)
            continue
        labeled = _labeled_bullet(line)
        if labeled:
            pairs = [list(labeled)]
            i += 1
            while i < len(lines):
                nxt = _labeled_bullet(lines[i])
                if not nxt:
                    break
                pairs.append(list(nxt))
                i += 1
            if len(pairs) >= 2:
                _add_table(doc, ("Item", "Basis"), pairs)
            else:
                _add_paragraph(doc, f"• {pairs[0][0]}: {pairs[0][1]}", size_pt=10, space_after=3)
            continue
        bullet = re.match(r"^(\s*)([-*•]|\d+\.)\s+(.*)$", line.strip())
        if bullet:
            _add_paragraph(doc, f"• {bullet.group(3)}", size_pt=10, space_after=2)
            i += 1
            continue
        _add_paragraph(doc, line, size_pt=10)
        i += 1


def _append_structured_report(doc, result: Mapping[str, Any], markdown: str) -> None:
    decoded = _decoded_rows(result)
    options = _eval_options(result)
    matrix = [row for row in (result.get("evaluation_matrix") or []) if isinstance(row, Mapping)]
    item_noun = _item_noun(result)
    selected = str(result.get("selected_model") or "").strip()
    n = 1
    implication_label = "Mixing implication" if item_noun == "mixing" else "Implication"

    basis_md = _markdown_section(markdown, "design basis from dir")
    if decoded or _extract_md_table(basis_md):
        _add_paragraph(doc, f"{n}. Design basis from DIR code", size_pt=13, bold=True, color=NAVY, space_before=12, space_after=4)
        n += 1
        intro = _section_intro(basis_md)
        if intro:
            _add_paragraph(doc, intro, size_pt=10)
        basis_rows = []
        unknown_any = False
        for row in decoded:
            unknown = bool(row.get("unknown"))
            unknown_any = unknown_any or unknown
            selected_basis = str(row.get("option_text") or row.get("value") or "")
            implication = str(row.get("implication") or row.get("mixing_implication") or "")
            if unknown and not implication:
                implication = (
                    "Assume the most likely industrial case for this duty; "
                    "confirm when project data exist."
                )
            basis_rows.append([str(row.get("label") or ""), selected_basis, implication])
        if unknown_any and not intro:
            _add_paragraph(
                doc,
                "Unknown / TBD selections are not yet defined by project engineering. "
                "The evaluation assumes the most likely case and states consequences.",
                size_pt=10,
            )
        _add_md_table_or(
            doc,
            basis_md,
            ("DIR element", "Selected basis", implication_label),
            basis_rows,
        )

    objectives = _as_str_list(result.get("objectives"))
    failures = _as_str_list(result.get("failure_modes"))
    obj_md = _markdown_section(markdown, "objectives and failure")
    obj_rows = _objective_rows(obj_md, objectives, result)
    if obj_rows or failures:
        _add_paragraph(
            doc,
            f"{n}. {item_noun.title()} objectives and failure modes",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        n += 1
        if obj_rows:
            _add_table(doc, ("Step", "Objective", "Key control"), obj_rows)
        if failures:
            _add_paragraph(doc, "; ".join(failures), size_pt=10, label="Failure modes: ")

    practice_md = _markdown_section(markdown, "industry best practice")
    if practice_md:
        _add_paragraph(
            doc,
            f"{n}. Industry best practice for this duty",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        n += 1
        intro = _section_intro(practice_md)
        if intro:
            _add_paragraph(doc, intro, size_pt=10)

    short_md = _markdown_section(markdown, "best-fit", "strong-fit", "system shortlist")
    if options or matrix or _extract_md_table(short_md):
        _add_paragraph(
            doc,
            f"{n}. Best-fit {item_noun}-system shortlist",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        n += 1
        intro = _section_intro(short_md)
        if intro:
            _add_paragraph(doc, intro, size_pt=10)
        short_rows = []
        by_name = {str(row.get("option") or "").strip().lower(): row for row in matrix}
        for opt in options:
            name = str(opt.get("name") or "Option")
            fit = str(opt.get("fit") or "").replace("_", "-")
            mrow = by_name.get(name.lower(), {})
            short_rows.append(
                [
                    name,
                    _rating_dots(mrow.get("technical_fit") or fit),
                    _rating_dots(mrow.get("gmp") or ""),
                    _rating_dots(mrow.get("scale_up_risk") or ""),
                    str(fit or "").replace("-", " ").title() or "—",
                    _fit_position(fit),
                ]
            )
        if not short_rows:
            for row in matrix:
                short_rows.append(
                    [
                        str(row.get("option") or ""),
                        _rating_dots(row.get("technical_fit") or ""),
                        _rating_dots(row.get("gmp") or ""),
                        _rating_dots(row.get("scale_up_risk") or ""),
                        str(row.get("reliability") or row.get("rank") or ""),
                        "",
                    ]
                )
        _add_md_table_or(
            doc,
            short_md,
            ("Option", "Technical fit", "GMP / hygienic", "Scale-up risk", "Fit", "Position"),
            short_rows,
        )

    if options:
        _add_paragraph(doc, f"{n}. Option evaluation", size_pt=13, bold=True, color=NAVY, space_before=12, space_after=4)
        n += 1
        for idx, opt in enumerate(options, start=1):
            name = str(opt.get("name") or f"Option {idx}")
            _add_paragraph(doc, f"{idx}. {name}", size_pt=11, bold=True, color=NAVY, space_before=8, space_after=3)
            apps = _as_str_list(opt.get("industrial_applications"))
            if apps:
                _add_paragraph(doc, "; ".join(apps), size_pt=10, label="Industrial applications: ")
            rows = []
            pros = _as_str_list(opt.get("pros"))
            cons = _as_str_list(opt.get("cons"))
            mfrs = _as_str_list(opt.get("manufacturers"))
            if pros:
                rows.append(["Pros", "\n".join(f"• {p}" for p in pros)])
            if cons:
                rows.append(["Cons / watchouts", "\n".join(f"• {c}" for c in cons)])
            if mfrs:
                rows.append(["Manufacturers", "; ".join(mfrs)])
            if rows:
                _add_label_table(doc, rows)

    rec_basis = str(result.get("recommended_basis") or selected or "").strip()
    specs = _as_str_list(result.get("preliminary_specs"))
    spec_md = _markdown_section(
        markdown, "recommended preliminary", "preliminary specification", "preliminary basis"
    )
    if rec_basis or specs or _extract_md_table(spec_md):
        _add_paragraph(
            doc,
            f"{n}. Recommended preliminary basis of design",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        n += 1
        if rec_basis:
            _add_callout(doc, "", f"Selected basis: {rec_basis}")
        intro = _section_intro(spec_md)
        if intro:
            _add_paragraph(doc, intro, size_pt=10)
        alt = str(result.get("alternate_basis") or "").strip()
        if alt:
            _add_paragraph(doc, alt, size_pt=10, label="Alternate: ")
        spec_rows = []
        for spec in specs:
            key, val = _split_labeled(spec)
            spec_rows.append([key, val or spec])
        _add_md_table_or(doc, spec_md, ("Specification item", "Preliminary basis"), spec_rows)

    recipe = _markdown_section(markdown, "operating recipe", "qualification")
    if recipe:
        _add_paragraph(
            doc,
            f"{n}. Suggested operating recipe for qualification",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        n += 1
        for line in recipe.splitlines():
            stripped = line.strip()
            if stripped and not stripped.startswith("|") and not _heading_text(stripped):
                _add_paragraph(doc, f"• {stripped.lstrip('-* ')}", size_pt=10, space_after=2)

    excluded_raw = result.get("do_not_specify")
    excluded = excluded_raw if isinstance(excluded_raw, list) else _as_str_list(excluded_raw)
    excl_md = _markdown_section(markdown, "not recommended", "do not specify")
    excl_rows = _exclusion_rows(excl_md, excluded)
    if excl_rows:
        _add_paragraph(
            doc,
            f"{n}. Options not recommended as primary basis",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        n += 1
        _add_table(doc, ("Technology", "Reason not primary for this DIR"), excl_rows)

    matrix_md = _markdown_section(markdown, "evaluation matrix", "option evaluation matrix")
    if matrix or _extract_md_table(matrix_md):
        _add_paragraph(
            doc,
            f"{n}. Preliminary option evaluation matrix",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        n += 1
        _add_md_table_or(
            doc,
            matrix_md,
            ("Option", "Technical fit", "GMP / cleanability", "Scale-up risk", "Cost / schedule", "Overall"),
            [
                [
                    str(row.get("option") or ""),
                    str(row.get("technical_fit") or ""),
                    str(row.get("gmp") or ""),
                    str(row.get("scale_up_risk") or ""),
                    str(row.get("cost_schedule") or ""),
                    str(row.get("reliability") or row.get("rank") or ""),
                ]
                for row in matrix
            ],
        )

    vendors = _as_str_list(result.get("manufacturers"))
    vendor_md = _markdown_section(markdown, "vendor / manufacturer", "manufacturer shortlist")
    if vendors or _extract_md_table(vendor_md):
        _add_paragraph(
            doc,
            f"{n}. Vendor / manufacturer shortlist",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        n += 1
        vendor_rows = []
        for item in vendors:
            left, right = _split_labeled(item)
            if right:
                vendor_rows.append([left, right])
            else:
                vendor_rows.append(["Manufacturer shortlist", item])
        if not vendor_rows:
            vendor_rows = [["Manufacturer shortlist", "; ".join(vendors)]]
        _add_md_table_or(doc, vendor_md, ("Role", "Likely manufacturers / suppliers"), vendor_rows)

    refs = _markdown_section(markdown, "references reviewed", "references")
    if refs:
        _add_paragraph(
            doc,
            f"{n}. References reviewed",
            size_pt=13,
            bold=True,
            color=NAVY,
            space_before=12,
            space_after=4,
        )
        for line in refs.splitlines():
            if line.strip() and not line.strip().startswith("|"):
                _add_paragraph(doc, f"• {line.lstrip('-* ')}", size_pt=10, space_after=2)


def write_evaluation_report_docx(
    result: Mapping[str, Any],
    output_path: Path | str,
    *,
    title: str | None = None,
) -> Path:
    """Render the evaluation as an editable Word report."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    out = Path(output_path)
    if out.suffix.lower() != ".docx":
        out = out.with_suffix(".docx")
    out.parent.mkdir(parents=True, exist_ok=True)

    markdown = (result.get("datasheet_markdown") or "").strip() or _synthesize_markdown(result)
    display_title = _display_title(result, title)
    system = display_title
    dir_code = str(result.get("dir_code") or "").strip()
    application = str(result.get("application") or "").strip()
    item_noun = _item_noun(result)
    if str(result.get("schema_version") or "") == "equipment_sizing_v1":
        subtitle = f"Preliminary {item_noun} sizing and recommended basis of design"
    else:
        subtitle = f"Preliminary {item_noun}-system option evaluation and recommended basis of design"
    header_title = display_title
    if dir_code:
        header_title += f" • Basis: user DIR code {dir_code}"
    recommendation = _extract_recommendation(result, markdown)

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.7)
    _set_header_footer(doc, header_title)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title_p.add_run(display_title)
    _set_run(title_run, size_pt=20, bold=True, color=NAVY)
    title_p.paragraph_format.space_after = Pt(2)

    _add_paragraph(doc, subtitle, size_pt=10, color=MUTED, space_after=8)

    attr_pairs = _key_spec_pairs(result)
    if attr_pairs:
        chips = doc.add_table(rows=2, cols=len(attr_pairs))
        chips.style = "Table Grid"
        for i, (label, value) in enumerate(attr_pairs):
            _set_cell_text(chips.rows[0].cells[i], label, header=True)
            _set_cell_text(chips.rows[1].cells[i], value)

    meta = [c for c in (
        f"Validated DIR: {dir_code}" if dir_code else "",
        f"Application assumption: {application}" if application else "",
    ) if c]
    if meta:
        meta_table = doc.add_table(rows=1, cols=len(meta))
        meta_table.style = "Table Grid"
        for i, cell_text in enumerate(meta):
            _shade(meta_table.rows[0].cells[i], CALLOUT_BG)
            _set_cell_text(meta_table.rows[0].cells[i], cell_text)
            _shade(meta_table.rows[0].cells[i], CALLOUT_BG)
            if meta_table.rows[0].cells[i].paragraphs[0].runs:
                _set_run(meta_table.rows[0].cells[i].paragraphs[0].runs[0], size_pt=9, bold=True, color=NAVY)
        doc.add_paragraph("")

    if recommendation:
        _add_callout(doc, "Recommendation in one line", recommendation)

    if _has_structured_eval(result):
        _append_structured_report(doc, result, markdown)
    else:
        _append_unstructured_markdown(doc, markdown, system, display_title)

    try:
        doc.save(str(out))
        return out
    except PermissionError:
        from datetime import datetime

        stamped = out.with_name(f"{out.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{out.suffix}")
        doc.save(str(stamped))
        return stamped


def build_evaluation_docx(
    result: Mapping[str, Any],
    *,
    output_path: Path | str,
    title: str | None = None,
) -> Path:
    """Public SDK entry used by evaluator and sizing templates."""
    return write_evaluation_report_docx(result, output_path, title=title)


def evaluation_docx_text(path: Path | str) -> str:
    """Extract paragraph and table text for tests."""
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)
