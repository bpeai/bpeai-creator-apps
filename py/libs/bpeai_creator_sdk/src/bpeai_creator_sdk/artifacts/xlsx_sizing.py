"""Generic equipment-sizing Excel workbook (no per-equipment formula engine)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Mapping, Sequence

from .report_content import _is_table_divider, _split_table_row


def parse_markdown_table(markdown: str) -> tuple[list[str], list[list[str]]]:
    """Return (headers, rows) for the first markdown table in ``markdown``."""
    lines = (markdown or "").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.strip().startswith("|") and i + 1 < len(lines) and _is_table_divider(lines[i + 1]):
            headers = _split_table_row(line)
            rows: list[list[str]] = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not _is_table_divider(lines[i]):
                    rows.append(_split_table_row(lines[i]))
                i += 1
            if headers:
                return headers, rows
        i += 1
    return [], []


def _looks_like_excel_formula(text: str) -> bool:
    value = str(text or "").strip()
    return bool(value) and value.startswith("=")


def _write_result_cell(cell, raw: str) -> None:
    text = str(raw or "").strip()
    if _looks_like_excel_formula(text):
        cell.value = text
        return
    compact = text.replace(",", "")
    try:
        if re.fullmatch(r"-?\d+(\.\d+)?([eE][+-]?\d+)?", compact):
            cell.value = float(compact) if "." in compact or "e" in compact.lower() else int(compact)
            return
    except ValueError:
        pass
    cell.value = text


def _capacity_rows(result: Mapping[str, Any]) -> list[tuple[str, str]]:
    cap = result.get("capacity") if isinstance(result.get("capacity"), Mapping) else {}
    rows = [
        ("System name", str(result.get("system_name") or result.get("equipment_name") or "")),
        ("Application", str(result.get("application") or "")),
        ("DIR code", str(result.get("dir_code") or "")),
        ("Sized item", str(result.get("sized_item") or "")),
        ("Capacity value", str(cap.get("value") or "")),
        ("Capacity unit", str(cap.get("unit") or "")),
        ("Capacity basis", str(cap.get("basis") or "")),
    ]
    return [(k, v) for k, v in rows if str(v).strip()]


def build_sizing_xlsx(
    result: Mapping[str, Any],
    *,
    output_path: Path | str,
) -> Path:
    """Write Inputs / Calculations / Summary / Audit sheets from sizing JSON.

    Calculation Result cells that start with ``=`` are stored as Excel formulas.
    Domain math stays in the knowledge pack / LLM table — this helper does not
    embed agitator- or pump-specific equations.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - optional local dep
        raise RuntimeError("openpyxl is required to write the sizing workbook") from exc

    out = Path(output_path)
    if out.suffix.lower() != ".xlsx":
        out = out.with_suffix(".xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)

    wb = Workbook()

    inputs = wb.active
    inputs.title = "Inputs"
    inputs.append(["Parameter", "Value"])
    for cell in inputs[1]:
        cell.fill = header_fill
        cell.font = header_font
    for key, value in _capacity_rows(result):
        inputs.append([key, value])
    for assumption in result.get("assumptions") or []:
        text = str(assumption or "").strip()
        if text:
            inputs.append(["Assumption", text])
    for used in result.get("inputs_used") or []:
        text = str(used or "").strip()
        if text:
            inputs.append(["Input used", text])

    calc = wb.create_sheet("Calculations")
    table_md = str(result.get("excel_ready_table") or "").strip()
    if not table_md:
        table_md = str(result.get("datasheet_markdown") or "")
    headers, rows = parse_markdown_table(table_md)
    if not headers:
        headers = ["Item", "Method/formula", "Result", "Unit", "Basis"]
    calc.append(headers)
    for cell in calc[1]:
        cell.fill = header_fill
        cell.font = header_font
    result_idx = next(
        (i for i, h in enumerate(headers) if "result" in str(h).lower()),
        2 if len(headers) > 2 else len(headers) - 1,
    )
    for row in rows:
        padded = list(row) + [""] * max(0, len(headers) - len(row))
        excel_row = []
        for i, value in enumerate(padded[: len(headers)]):
            excel_row.append(value)
        calc.append(excel_row)
        written = calc[calc.max_row]
        if 0 <= result_idx < len(written):
            _write_result_cell(written[result_idx], padded[result_idx])

    summary = wb.create_sheet("Summary")
    summary.append(["Item", "Value", "Unit"])
    for cell in summary[1]:
        cell.fill = header_fill
        cell.font = header_font
    specs = result.get("key_specs") or []
    if isinstance(specs, Sequence):
        for spec in specs:
            if not isinstance(spec, Mapping):
                continue
            key = str(spec.get("key") or "").strip()
            if not key:
                continue
            summary.append(
                [
                    key,
                    spec.get("value") if spec.get("value") is not None else "",
                    spec.get("unit") if spec.get("unit") not in (None, "") else "",
                ]
            )
    selected = str(result.get("selected_model") or "").strip()
    if selected:
        summary.append(["Sized-item concept", selected, ""])

    audit = wb.create_sheet("Audit")
    audit.append(["Kind", "Note"])
    for cell in audit[1]:
        cell.fill = header_fill
        cell.font = header_font
    for missing in result.get("missing_inputs") or []:
        audit.append(["Missing input", str(missing)])
    for basis in result.get("source_basis") or []:
        audit.append(["Source basis", str(basis)])
    for warning in result.get("sme_warnings") or []:
        audit.append(["Warning", str(warning)])
    if audit.max_row == 1:
        audit.append(["Note", "Preliminary sizing; vendor confirmation required."])

    for sheet in wb.worksheets:
        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            width = 12
            for cell in column:
                width = max(width, min(len(str(cell.value or "")), 48))
            sheet.column_dimensions[letter].width = width + 2

    try:
        wb.save(str(out))
        return out
    except PermissionError:
        from datetime import datetime

        stamped = out.with_name(f"{out.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{out.suffix}")
        wb.save(str(stamped))
        return stamped
